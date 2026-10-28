"""Turn Travel Clean.xlsx (+ an optional Google Earth .kml) into supabase/seed.sql.

usage: python3 export_seed.py "Travel Clean.xlsx" Nathan.kml seed.sql [data.json] [--email you@example.com]
The xlsx must have its formula values saved (open and save it once in Excel/Sheets/LibreOffice).
"""
import sys, re, json, math, datetime, xml.etree.ElementTree as ET
import openpyxl

args = [a for a in sys.argv[1:] if not a.startswith('--')]
email = sys.argv[sys.argv.index('--email') + 1] if '--email' in sys.argv else ''
if email in args: args.remove(email)
xlsx, kml, out = args[0], args[1], args[2]
json_out = args[3] if len(args) > 3 else None
ME = 'nate'

wb = openpyxl.load_workbook(xlsx, data_only=True)
def rows(tab):
    ws = wb[tab]; head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] not in (None, '')]
def d(v): return v.date().isoformat() if isinstance(v, datetime.datetime) else (v.isoformat() if isinstance(v, datetime.date) else None)
def s(v): return None if v in (None, '') else str(v).strip()
def n(v):
    try: return None if v in (None, '') else float(v)
    except ValueError: return None

def hav(a, b):
    R = 3958.8; r = math.radians
    h = math.sin(r(b[0]-a[0])/2)**2 + math.cos(r(a[0]))*math.cos(r(b[0]))*math.sin(r(b[1]-a[1])/2)**2
    return 2*R*math.asin(math.sqrt(h))
def ll(v):
    try: a, b = [float(x) for x in str(v).split(',')]; return (a, b)
    except Exception: return None

GROUP = re.compile(r'team|family|roomates|roommates|church', re.I)
def slug(name): return re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-')[:40]

data = {k: [] for k in ['people','homes','countries','places','airports','trips','legs','stays','packing','layers','pins']}
people = {ME: dict(id=ME, name='Nate', kind='family', email=email or None, color='#2f7f8f')}

for c in rows('Countries'): data['countries'].append(dict(country=c['Country'], continent=s(c['Continent'])))
for p in rows('Places'): data['places'].append(dict(place=p['Place'], city=s(p['City']), state=s(p['State']), country=s(p['Country']), latlong=s(p['LatLong'])))
for a in rows('Airports'): data['airports'].append(dict(code=a['Code'], name=s(a['Airport']), place=s(a['Place']), latlong=s(a['LatLong'])))
for h in rows('Homes'):
    if d(h['From']) and h['Place']: data['homes'].append(dict(id='h-%s' % d(h['From']), person_id=ME, from_date=d(h['From']), place=h['Place']))
for t in rows('Trips'):
    data['trips'].append(dict(id=str(t['Trip ID']), name=s(t['Trip Name']), place=s(t['Place']), start_date=d(t['Start Date']), end_date=d(t['End Date']),
        date_confidence=s(t['Date Confidence']) or 'Exact', trip_type='Missions' if t['Missions'] == 'Yes' else None, status='Done',
        organization=s(t['Organization']), notes=s(t['Notes']), map_group=s(t['Map Group']), photo=s(t['Photo'])))
for l in rows('Legs'):
    who = [ME]
    for name in str(l['Traveled With'] or '').split(','):
        name = name.strip()
        if not name or name == 'Solo' or re.fullmatch(r'[0-9a-f]{8}', name): continue
        pid = slug(name)
        people.setdefault(pid, dict(id=pid, name=name, kind='group' if GROUP.search(name) else 'person'))
        if pid not in who: who.append(pid)
    lid = str(l['Leg ID'])
    data['legs'].append(dict(id=lid, trip_id=s(l['Trip ID']), mode=l['Mode'], date=d(l['Date']), end_date=d(l['End Date']), same_day=l['Same-Day Outing'] == 'Yes',
        from_text=s(l['From']), from_place=s(l['From Place']), to_text=s(l['To']), to_place=s(l['To Place']),
        from_airport=s(l['From Airport']), to_airport=s(l['To Airport']), miles=n(l['Miles']), people=who,
        date_confidence=s(l['Date Confidence']), notes=s(l['Notes']), photo=s(l['Photo']), name=s(l['Leg Name']), map_group=s(l['Map Group']),
        **{'class': 'First' if lid in ('F8', 'F166') else None}))
for st in rows('Stays'):
    data['stays'].append(dict(id=str(st['Stay ID']), trip_id=s(st['Trip ID']), place=st['Place'], arrive=d(st['Arrive']), leave=d(st['Leave']),
        people=[ME], confidence=s(st['Confidence']), notes=s(st['Notes'])))
data['people'] = list(people.values())

PACK = [('Any','Passport','Documents'),('Any','Phone charger','Tech'),('Any','Toothbrush','Toiletries'),('Any','Meds','Toiletries'),
        ('Any','Wallet and cards','Documents'),('Any','Headphones','Tech'),('Missions','Bible','Other'),('Missions','Ministry supplies','Other'),
        ('Missions','Modest clothes','Clothes'),('Work','Laptop','Tech'),('Fun','Camera','Tech'),('Fun','Swimsuit','Clothes'),('Family','Snacks','Other')]
data['packing'] = [dict(id='p%02d' % i, trip_type=a, item=b, category=c) for i, (a, b, c) in enumerate(PACK)]

# ---------------------------------------------------------------- Google Earth
def rdp(pts, eps):
    if len(pts) < 3: return pts
    a, b = pts[0], pts[-1]; dx, dy = b[0]-a[0], b[1]-a[1]; L = math.hypot(dx, dy) or 1e-12
    dmax, idx = 0, 0
    for i in range(1, len(pts)-1):
        dd = abs(dy*pts[i][0] - dx*pts[i][1] + b[0]*a[1] - b[1]*a[0]) / L
        if dd > dmax: dmax, idx = dd, i
    if dmax <= eps: return [a, b]
    return rdp(pts[:idx+1], eps)[:-1] + rdp(pts[idx:], eps)

if kml and kml != '-':
    ns = {'k': 'http://www.opengis.net/kml/2.2'}
    doc = ET.parse(kml).getroot().find('k:Document', ns)
    feats = []
    def walk(el, path):
        for ch in el:
            tag = ch.tag.split('}')[1]
            if tag == 'Folder': walk(ch, path + [(ch.findtext('k:name', default='', namespaces=ns) or '').strip()])
            elif tag == 'Placemark':
                props = dict(name=(ch.findtext('k:name', default='', namespaces=ns) or '').strip(), folder=path[0] if path else '',
                             sub=path[1] if len(path) > 1 else '', description=re.sub(r'<[^>]+>', ' ', ch.findtext('k:description', default='', namespaces=ns) or '').strip()[:500])
                pt = ch.find('.//k:Point/k:coordinates', ns)
                if pt is not None:
                    lo, la = [float(v) for v in pt.text.strip().split(',')[:2]]
                    feats.append(dict(type='Feature', properties=props, geometry=dict(type='Point', coordinates=[round(lo, 5), round(la, 5)])))
                for ls in ch.findall('.//k:LineString/k:coordinates', ns):
                    pts = [[float(x) for x in c.split(',')[:2]] for c in ls.text.split()]
                    pts = [[round(x, 4), round(y, 4)] for x, y in rdp(pts, 0.005)]
                    feats.append(dict(type='Feature', properties=props, geometry=dict(type='LineString', coordinates=pts)))
    walk(doc, [])
    layer = dict(id='earth1', name=re.sub(r'\.kml$', '', kml.split('/')[-1]).split('-')[-1], geojson=dict(type='FeatureCollection', features=feats))
    data['layers'].append(layer)
    # match pins to places (≤50 km) and to trips that touched that place
    P = [(p['place'], ll(p['latlong'])) for p in data['places'] if ll(p['latlong'])]
    touched = {}
    for x in data['stays']: touched.setdefault(x['place'], set()).add(x['trip_id'])
    for x in data['legs']:
        for k in ('from_place', 'to_place'): touched.setdefault(x[k], set()).add(x['trip_id'])
    trips = {t['id']: t for t in data['trips']}
    words = lambda v: {w for w in re.sub(r'[^a-z0-9 ]', ' ', (v or '').lower()).split() if len(w) > 2 and w not in ('the', 'and', 'trip', 'from', 'airport')}
    for i, f in enumerate(x for x in feats if x['geometry']['type'] == 'Point'):
        lo, la = f['geometry']['coordinates']; pr = f['properties']
        best = min(P, key=lambda q: hav((la, lo), q[1]))
        km = hav((la, lo), best[1]) * 1.609
        place = best[0] if km <= 50 else None
        trip = None
        if place and pr['folder'] != 'Janessa':        # that folder is someone else's travel
            cands = [trips[t] for t in touched.get(place, ()) if t in trips]
            score = lambda t: (2 if t['map_group'] and t['map_group'] == pr['folder'] else 0) + len(words(t['name'] + ' ' + (t['place'] or '')) & words(pr['sub']))
            cands.sort(key=score, reverse=True)
            if len(cands) == 1 or (len(cands) > 1 and score(cands[0]) > score(cands[1])): trip = cands[0]['id']
        data['pins'].append(dict(id='pin%03d' % i, layer_id='earth1', name=pr['name'], folder=pr['folder'], subfolder=pr['sub'],
                                 latlong='%.5f, %.5f' % (la, lo), notes=pr['description'] or None, place=place, trip_id=trip))

# ---------------------------------------------------------------- write SQL
def lit(v):
    if v is None: return 'null'
    if isinstance(v, bool): return 'true' if v else 'false'
    if isinstance(v, (int, float)): return repr(int(v)) if float(v).is_integer() else repr(v)
    if isinstance(v, list): return 'array[%s]::text[]' % ','.join(lit(x) for x in v) if v else "'{}'::text[]"
    if isinstance(v, dict): return lit(json.dumps(v, separators=(',', ':'))) + '::jsonb'
    return "'" + str(v).replace("'", "''") + "'"
order = ['countries', 'places', 'airports', 'people', 'homes', 'trips', 'legs', 'stays', 'packing', 'layers', 'pins']
def write(path, tables, note):
    with open(path, 'w') as f:
        f.write('-- Generated by tools/export_seed.py. %s\nbegin;\n' % note)
        for t in tables:
            for r in data[t]:
                cols = list(r.keys())
                f.write('insert into %s (%s) values (%s) on conflict do nothing;\n' % (t, ','.join('"%s"' % c for c in cols), ','.join(lit(r[c]) for c in cols)))
        f.write('commit;\n')
write(out, order[:-2], 'Run after schema.sql.')
if data['layers']: write(out.replace('.sql', '-google-earth.sql'), order[-2:], 'Run after seed.sql.')
if json_out: json.dump(data, open(json_out, 'w'))
print({k: len(v) for k, v in data.items()}, 'pins matched to places:', sum(1 for p in data['pins'] if p['place']), 'to trips:', sum(1 for p in data['pins'] if p['trip_id']))
